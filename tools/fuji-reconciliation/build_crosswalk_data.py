from __future__ import annotations

import argparse
import difflib
import json
import os
import re
import subprocess
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def ascii_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return text.encode("ascii", "ignore").decode("ascii")


def clean(value: object) -> str:
    text = ascii_text(value).upper().replace("?", "")
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def source_parts(name: str) -> tuple[str, str]:
    parts = name.split(",", 1)
    return clean(parts[0]), clean(parts[1] if len(parts) > 1 else "")


def db_display(last_name: str, first_name: str) -> str:
    return f"{last_name}, {first_name}".strip(" ,")


def name_score(source_name: str, candidate_name: str) -> tuple[float, float, float, bool]:
    source_last, source_given = source_parts(source_name)
    candidate_last, candidate_given = source_parts(candidate_name)
    last_score = difflib.SequenceMatcher(None, source_last, candidate_last).ratio()
    given_score = difflib.SequenceMatcher(None, source_given, candidate_given).ratio()
    source_first = source_given.split(" ", 1)[0] if source_given else ""
    candidate_first = candidate_given.split(" ", 1)[0] if candidate_given else ""
    first_token_match = bool(source_first and candidate_first and source_first == candidate_first)
    total = (last_score * 0.58) + (given_score * 0.42)
    if first_token_match:
        total = min(1.0, total + 0.04)
    return total, last_score, given_score, first_token_match


def number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def read_dtr(dtr_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(dtr_path, data_only=True, read_only=True)
    worksheet = workbook.active
    rows: list[dict[str, object]] = []
    for excel_row, values in enumerate(
        worksheet.iter_rows(min_col=1, max_col=71, values_only=True), start=1
    ):
        sequence, source_id, name = values[0], values[1], values[2]
        if not isinstance(sequence, (int, float)) or source_id is None or not isinstance(name, str):
            continue
        if not name.strip():
            continue
        rows.append(
            {
                "excel_row": excel_row,
                "sequence": int(sequence),
                "source_id": str(source_id),
                "source_name": name.strip(),
                "area": str(values[3] or "").strip(),
                "date_hired": values[4].isoformat() if hasattr(values[4], "isoformat") else str(values[4] or ""),
                "days_worked": number(values[20]),
                "regular_hours": number(values[22]),
                "late_minutes": number(values[23]) + number(values[35]) + number(values[43]) + number(values[51]) + number(values[60]),
                "overtime_hours": number(values[24]),
                "night_diff_hours": number(values[25]),
                "night_diff_ot_hours": number(values[26]),
                "undertime_minutes": number(values[27]) + number(values[39]) + number(values[47]) + number(values[55]) + number(values[64]),
                "meal_allowance": number(values[28]),
                "perfect_attendance": number(values[29]),
                "quarterly_attendance": number(values[30]),
                "hmo_deduction": number(values[31]),
                "rest_day_hours": number(values[34]),
                "rest_day_ot_hours": number(values[36]),
                "rest_day_night_diff_hours": number(values[37]),
                "rest_day_nd_ot_hours": number(values[38]),
                "regular_holiday_hours": number(values[42]),
                "regular_holiday_ot_hours": number(values[44]),
                "regular_holiday_night_diff_hours": number(values[45]),
                "regular_holiday_nd_ot_hours": number(values[46]),
                "special_holiday_hours": number(values[50]),
                "special_holiday_ot_hours": number(values[52]),
                "special_holiday_night_diff_hours": number(values[53]),
                "special_holiday_nd_ot_hours": number(values[54]),
                "sil_days": number(values[56]),
                "adjustment_amount": number(values[57]),
                "adjustment_hours": number(values[59]),
                "adjustment_ot_hours": number(values[61]),
                "adjustment_night_diff_hours": number(values[62]),
                "adjustment_nd_ot_hours": number(values[63]),
                "adjustment_perfect_attendance": number(values[65]),
                "adjustment_quarterly_attendance": number(values[66]),
                "adjustment_meal_allowance": number(values[67]),
                "adjustment_sil": number(values[68]),
                "paternity_leave": number(values[69]),
                "adjustment_legal_holiday": number(values[70]),
            }
        )
    workbook.close()
    return rows


def read_pdf_names(pdf_text_path: Path) -> list[str]:
    text = pdf_text_path.read_text(encoding="utf-8")
    names: list[str] = []
    for block in text.split("TAASCOR MANAGEMENT & GENERAL SERVICES CORPORATION"):
        match = re.search(r"Employee Name\s+(.+?)\s*\n", block)
        if match:
            names.append(match.group(1).strip())
    return names


def read_hris(
    mysql_bin: str,
    db_host: str,
    db_port: int,
    db_user: str,
    database: str,
    client_id: int,
) -> list[dict[str, object]]:
    query = (
        "SELECT e.employee_id,COALESCE(e.payroll_employee_id,''),"
        "COALESCE(e.first_name,''),COALESCE(e.last_name,''),"
        "COALESCE(s.daily_salary,0),COALESCE(e.status,'') "
        "FROM employee_list e LEFT JOIN employee_salary s ON e.employee_id=s.employee_id "
        f"WHERE e.client_id={int(client_id)}"
    )
    environment = os.environ.copy()
    if environment.get("DB_PASSWORD"):
        environment["MYSQL_PWD"] = environment["DB_PASSWORD"]
    result = subprocess.run(
        [
            mysql_bin,
            f"--host={db_host}",
            f"--port={db_port}",
            f"--user={db_user}",
            database,
            "--batch",
            "--raw",
            "--skip-column-names",
            f"--execute={query}",
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=environment,
    )
    rows: list[dict[str, object]] = []
    for line in result.stdout.splitlines():
        fields = line.split("\t")
        rows.append(
            {
                "employee_id": int(fields[0]),
                "payroll_employee_id": fields[1],
                "first_name": fields[2],
                "last_name": fields[3],
                "daily_salary": float(fields[4]),
                "status": fields[5],
                "display_name": db_display(fields[3], fields[2]),
            }
        )
    return rows


def match_pdf(dtr_rows: list[dict[str, object]], pdf_names: list[str]) -> None:
    unused = set(range(len(pdf_names)))
    for row in dtr_rows:
        scored = []
        for index in unused:
            total, last_score, given_score, first_match = name_score(str(row["source_name"]), pdf_names[index])
            if last_score >= 0.65:
                scored.append((total, last_score, given_score, first_match, index))
        scored.sort(reverse=True)
        if not scored:
            row.update({"pdf_name": "", "pdf_match_score": 0.0, "pdf_match_status": "Not matched"})
            continue
        total, last_score, given_score, first_match, index = scored[0]
        if total >= 0.86 and last_score >= 0.85:
            status = "Matched - high confidence"
        elif total >= 0.72 and last_score >= 0.75:
            status = "Matched - review"
        else:
            status = "Not matched"
        if status != "Not matched":
            unused.remove(index)
        row.update(
            {
                "pdf_name": pdf_names[index],
                "pdf_match_score": round(total, 4),
                "pdf_match_status": status,
            }
        )


def match_hris(dtr_rows: list[dict[str, object]], hris_rows: list[dict[str, object]]) -> None:
    for row in dtr_rows:
        scored = []
        for candidate in hris_rows:
            total, last_score, given_score, first_match = name_score(
                str(row["source_name"]), str(candidate["display_name"])
            )
            if last_score < 0.45:
                continue
            active_bonus = 0.01 if clean(candidate["status"]) == "ACTIVE" else 0.0
            scored.append((min(1.0, total + active_bonus), last_score, given_score, first_match, candidate))
        scored.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
        best = scored[0] if scored else (0.0, 0.0, 0.0, False, None)
        second = scored[1] if len(scored) > 1 else (0.0, 0.0, 0.0, False, None)
        score, last_score, given_score, first_match, candidate = best
        gap = score - second[0]
        candidate_active = bool(candidate and clean(candidate["status"]) == "ACTIVE")
        candidate_salary_ready = bool(candidate and float(candidate["daily_salary"]) > 0)
        high_confidence = bool(
            candidate
            and candidate_active
            and score >= 0.86
            and last_score >= 0.82
            and (first_match or given_score >= 0.82)
            and gap >= 0.025
        )
        match_status = "Suggested - high confidence" if high_confidence else "Owner review required"
        reasons = []
        if not candidate:
            reasons.append("No HRIS candidate")
        else:
            if last_score < 0.82:
                reasons.append("Surname differs")
            if not first_match and given_score < 0.82:
                reasons.append("Given name differs")
            if gap < 0.025:
                reasons.append("Competing HRIS candidate")
            if not candidate_active:
                reasons.append("Candidate is not active")
            if not candidate_salary_ready:
                reasons.append("Salary missing")
        if not candidate:
            payroll_readiness_status = "Blocked - no HRIS candidate"
        elif not candidate_active:
            payroll_readiness_status = "Blocked - inactive employee"
        elif not high_confidence:
            payroll_readiness_status = "Blocked - identity review"
        elif not candidate_salary_ready:
            payroll_readiness_status = "Blocked - salary missing"
        else:
            payroll_readiness_status = "Ready for earnings test"
        row.update(
            {
                "hris_match_status": match_status,
                "hris_match_score": round(score, 4),
                "hris_match_gap": round(gap, 4),
                "suggested_employee_id": candidate["employee_id"] if candidate else "",
                "suggested_payroll_employee_id": candidate["payroll_employee_id"] if candidate else "",
                "suggested_hris_name": candidate["display_name"] if candidate else "",
                "suggested_daily_salary": candidate["daily_salary"] if candidate else 0.0,
                "suggested_hris_status": candidate["status"] if candidate else "",
                "payroll_readiness_status": payroll_readiness_status,
                "second_candidate_employee_id": second[4]["employee_id"] if second[4] else "",
                "second_candidate_name": second[4]["display_name"] if second[4] else "",
                "second_candidate_score": round(second[0], 4),
                "review_reason": "; ".join(reasons) if reasons else "Verify source ID against owner crosswalk",
                "approval_status": "Pending owner review",
                "approved_employee_id": "",
                "reviewer_note": "",
            }
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a private Fuji DTR/PDF/HRIS crosswalk. Output contains payroll PII."
    )
    parser.add_argument("--dtr", required=True, type=Path, help="Fuji timekeeping workbook")
    parser.add_argument("--pdf-text", required=True, type=Path, help="Text extracted from the expected payslip PDF")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "private-output" / "crosswalk_data.json",
    )
    parser.add_argument("--mysql-bin", default=os.getenv("MYSQL_BIN", "mysql"))
    parser.add_argument("--db-host", default=os.getenv("DB_HOST", "127.0.0.1"))
    parser.add_argument("--db-port", type=int, default=int(os.getenv("DB_PORT", "3306")))
    parser.add_argument("--db-user", default=os.getenv("DB_USERNAME", "root"))
    parser.add_argument("--database", default=os.getenv("DB_DATABASE", "taascor_hris"))
    parser.add_argument("--client-id", required=True, type=int)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    dtr_rows = read_dtr(args.dtr)
    pdf_names = read_pdf_names(args.pdf_text)
    hris_rows = read_hris(
        args.mysql_bin,
        args.db_host,
        args.db_port,
        args.db_user,
        args.database,
        args.client_id,
    )
    match_pdf(dtr_rows, pdf_names)
    match_hris(dtr_rows, hris_rows)
    summary = {
        "dtr_rows": len(dtr_rows),
        "pdf_payslips": len(pdf_names),
        "pdf_matched_high_confidence": sum(
            1 for row in dtr_rows if row["pdf_match_status"] == "Matched - high confidence"
        ),
        "pdf_matched_review": sum(
            1 for row in dtr_rows if row["pdf_match_status"] == "Matched - review"
        ),
        "pdf_not_matched": sum(1 for row in dtr_rows if row["pdf_match_status"] == "Not matched"),
        "hris_high_confidence": sum(
            1 for row in dtr_rows if row["hris_match_status"] == "Suggested - high confidence"
        ),
        "hris_owner_review": sum(
            1 for row in dtr_rows if row["hris_match_status"] == "Owner review required"
        ),
        "hris_inactive_candidate": sum(
            1 for row in dtr_rows if row["payroll_readiness_status"] == "Blocked - inactive employee"
        ),
        "hris_salary_missing": sum(
            1 for row in dtr_rows if row["payroll_readiness_status"] == "Blocked - salary missing"
        ),
        "hris_payroll_ready": sum(
            1 for row in dtr_rows if row["payroll_readiness_status"] == "Ready for earnings test"
        ),
        "pdf_names_not_consumed": len(pdf_names)
        - sum(1 for row in dtr_rows if row["pdf_match_status"] != "Not matched"),
        "status_counts": dict(Counter(str(row["hris_match_status"]) for row in dtr_rows)),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"summary": summary, "rows": dtr_rows}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
